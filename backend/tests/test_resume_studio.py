from __future__ import annotations

import io
import zipfile
from uuid import UUID
from docx import Document

from conftest import register_verified, seed_profile
from app.core.database import SessionLocal
from app.models.entities import UploadedFile


def _docx_bytes() -> bytes:
    out = io.BytesIO(); doc = Document()
    doc.add_paragraph("TEST CANDIDATE")
    doc.add_paragraph("Senior Product Engineer")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Senior Product Engineer at Clearline since 2023")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("React, TypeScript, PostgreSQL")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("B.Tech in Computer Science, PES University, 2018")
    doc.save(out); return out.getvalue()


def test_studio_upload_is_private_versioned_and_idempotent(client):
    register_verified(client)
    data = _docx_bytes()
    first = client.post("/api/v1/resume-studio/resumes", files={"file": ("resume.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert first.status_code == 202, first.text
    body = first.json()["data"]
    assert body["version"] == 1
    assert body["status"] in {"queued", "completed"}
    assert "storage_key" not in body
    replay = client.post("/api/v1/resume-studio/resumes", files={"file": ("copy.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert replay.status_code == 202
    assert replay.json()["data"]["id"] == body["id"]
    assert replay.json()["data"]["idempotent_replay"] is True


def test_studio_rejects_malformed_and_cross_account_access(client):
    register_verified(client)
    empty = client.post("/api/v1/resume-studio/resumes", files={"file": ("empty.pdf", b"", "application/pdf")})
    assert empty.status_code == 415
    bad = client.post("/api/v1/resume-studio/resumes", files={"file": ("resume.pdf", b"not a pdf", "application/pdf")})
    assert bad.status_code == 415
    data = _docx_bytes()
    uploaded = client.post("/api/v1/resume-studio/resumes", files={"file": ("resume.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()["data"]
    client.post("/api/v1/auth/logout")
    assert client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}").status_code == 401
    register_verified(client, "other@example.com")
    assert client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}").status_code == 404
    assert client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}/facts").status_code == 404
    assert client.delete(f"/api/v1/resume-studio/resumes/{uploaded['id']}").status_code == 404


def test_resume_processing_can_retry_and_owner_can_delete(client):
    register_verified(client)
    data = _docx_bytes()
    uploaded = client.post("/api/v1/resume-studio/resumes", files={"file": ("resume.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()["data"]
    assert client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}").json()["data"]["status"] == "completed"
    with SessionLocal() as db:
        record = db.get(UploadedFile, UUID(uploaded["id"]))
        record.processing_status = "failed"
        record.extraction_error = "simulated_failure"
        db.commit()
    retried = client.post(f"/api/v1/resume-studio/resumes/{uploaded['id']}/retry")
    assert retried.status_code == 202, retried.text
    assert client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}").json()["data"]["status"] == "completed"
    deleted = client.delete(f"/api/v1/resume-studio/resumes/{uploaded['id']}")
    assert deleted.status_code == 200, deleted.text
    assert client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}").status_code == 404


def test_fact_review_suggestions_target_templates_and_exports(client):
    register_verified(client); seed_profile(client)
    data = _docx_bytes()
    uploaded = client.post("/api/v1/resume-studio/resumes", files={"file": ("resume.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()["data"]
    facts = client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}/facts")
    assert facts.status_code == 200
    if facts.json()["data"]:
        assert all(fact["status"] == "pending" for fact in facts.json()["data"])
        confirmed = client.post("/api/v1/profile/verify", json={"decisions": [{"entry_id": fact["id"], "action": "accept"} for fact in facts.json()["data"]]})
        assert confirmed.status_code == 200
        assert all(fact["status"] == "confirmed" for fact in client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}/facts").json()["data"])
    requested = client.post(f"/api/v1/resume-studio/resumes/{uploaded['id']}/suggestions")
    assert requested.status_code == 202
    suggestions = client.get(f"/api/v1/resume-studio/resumes/{uploaded['id']}/suggestions")
    assert suggestions.status_code == 200
    templates = client.get("/api/v1/resume-studio/templates").json()["data"]
    assert {x["key"] for x in templates} == {"ats", "modern", "minimal", "professional"}
    created_target = client.post("/api/v1/resume-studio/target", json={"title": "Staff Frontend Engineer", "company": "Northstar Labs", "description": "We need React, TypeScript, accessibility and GraphQL experience to lead frontend platform work and improve performance.", "document_type": "cover_letter", "idempotency_key": "resume-studio-job-target"})
    assert created_target.status_code == 202, created_target.text
    job_id = created_target.json()["data"]["job_id"]
    target = client.post("/api/v1/resume-studio/target", json={"job_id": job_id, "document_type": "resume", "template": "modern", "source_file_id": uploaded["id"], "idempotency_key": "resume-studio-test-key"})
    assert target.status_code == 202, target.text
    task = client.get(f"/api/v1/tasks/{target.json()['data']['task_id']}").json()["data"]
    assert task["status"] == "succeeded", task
    doc_id = task["result"]["document_id"]
    doc = client.get(f"/api/v1/documents/{doc_id}").json()["data"]
    assert doc["template_key"] == "modern"
    assert doc["source_file_id"] == uploaded["id"]
    assert doc["claim_report"]["unsupported_claims"] == 0
    assert client.post(f"/api/v1/documents/{doc_id}/approve").status_code == 200
    assert client.get(f"/api/v1/documents/{doc_id}/download?format=pdf").status_code == 200
    assert client.get(f"/api/v1/documents/{doc_id}/download?format=docx").status_code == 200
    replay = client.post("/api/v1/resume-studio/target", json={"job_id": job_id, "document_type": "resume", "template": "modern", "source_file_id": uploaded["id"], "idempotency_key": "resume-studio-test-key"})
    assert replay.json()["data"]["idempotent_replay"] is True


def test_privacy_export_includes_studio_records(client):
    register_verified(client)
    data = _docx_bytes()
    client.post("/api/v1/resume-studio/resumes", files={"file": ("resume.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    exported = client.get("/api/v1/privacy/export")
    assert exported.status_code == 200
    with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
        assert b"resume_uploads" in archive.read("careerpilot-data.json")
