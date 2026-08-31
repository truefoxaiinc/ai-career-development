import io
from docx import Document
from conftest import register_verified

def test_resume_upload_parse_and_candidate_verification(client):
    register_verified(client)
    doc=Document();doc.add_paragraph("Asha Verma");doc.add_paragraph("Senior React Engineer");doc.add_paragraph("React TypeScript PostgreSQL Accessibility");doc.add_paragraph("B.Tech in Computer Science, Example University");doc.add_paragraph("Senior Product Engineer at Clearline since 2023")
    buf=io.BytesIO();doc.save(buf)
    upload=client.post("/api/v1/profile/resume",files={"file":("resume.docx",buf.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert upload.status_code==202,upload.text
    task=client.get(f"/api/v1/tasks/{upload.json()['data']['task_id']}").json()["data"]
    assert task["status"]=="succeeded"
    extracted=client.get("/api/v1/profile/extracted").json()["data"]
    assert extracted and all(not x.get("verified",False) for x in extracted)
    decisions=[{"entry_id":x["id"],"action":"accept"} for x in extracted]
    verified=client.post("/api/v1/profile/verify",json={"decisions":decisions})
    assert verified.status_code==200
    profile=client.get("/api/v1/profile").json()["data"]
    assert any(e["verified"] and e["entry_type"]=="skill" for e in profile["entries"])
