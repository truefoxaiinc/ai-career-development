"""Polished, template-aware PDF and Word exports for career documents."""

from __future__ import annotations

import io
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate


PDF_ACCENTS = {
    "ats": "#111827",
    "modern": "#087e8b",
    "minimal": "#52525b",
    "professional": "#1d4ed8",
}
DOCX_ACCENTS = {
    "ats": RGBColor(17, 24, 39),
    "modern": RGBColor(8, 126, 139),
    "minimal": RGBColor(82, 82, 91),
    "professional": RGBColor(29, 78, 216),
}


def _lines(document):
    return [line.strip() for line in document.content.splitlines() if line.strip()]


def _bullet_text(line: str) -> str:
    for prefix in ("Ã¢â‚¬Â¢", "â€¢", "•", "-"):
        if line.startswith(prefix):
            return line[len(prefix):].strip()
    return line


def _is_bullet(line: str) -> bool:
    return line.startswith(("Ã¢â‚¬Â¢", "â€¢", "•", "-"))


def _is_heading(line: str) -> bool:
    known = {
        "PROFESSIONAL SUMMARY", "PROFESSIONAL EXPERIENCE", "EXPERIENCE",
        "KEY ACHIEVEMENTS", "SELECTED PROJECTS", "PROJECTS",
        "CORE SKILLS", "SKILLS", "EDUCATION", "CERTIFICATIONS AND LICENSES",
        "CERTIFICATIONS", "PUBLICATIONS", "LANGUAGES", "ADDITIONAL INFORMATION",
    }
    return line.upper() in known


def _resume_header(lines, story, styles, accent):
    if not lines:
        return 0
    story.append(Paragraph(escape(lines[0]), styles["name"]))
    if len(lines) > 1:
        story.append(Paragraph(escape(lines[1]), styles["contact"]))
    next_line = 2
    if len(lines) > 2 and not _is_heading(lines[2]):
        story.append(Paragraph(escape(lines[2]), styles["headline"]))
        next_line = 3
    story.append(HRFlowable(width="100%", thickness=1.2, color=accent, spaceBefore=2, spaceAfter=6))
    return next_line


def export_pdf_bytes(doc) -> bytes:
    output = io.BytesIO()
    template = doc.template_key or "ats"
    accent = HexColor(PDF_ACCENTS.get(template, PDF_ACCENTS["ats"]))
    sample = getSampleStyleSheet()
    styles = {
        "name": ParagraphStyle("ResumeName", parent=sample["Normal"], fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=HexColor("#111827"), spaceAfter=2),
        "contact": ParagraphStyle("ResumeContact", parent=sample["Normal"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=HexColor("#52525b"), spaceAfter=5),
        "headline": ParagraphStyle("ResumeHeadline", parent=sample["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=accent, spaceAfter=6),
        "heading": ParagraphStyle("ResumeSection", parent=sample["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=accent, spaceBefore=9, spaceAfter=4, keepWithNext=True),
        "body": ParagraphStyle("ResumeBody", parent=sample["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13.4, textColor=HexColor("#27272a"), alignment=TA_LEFT, spaceAfter=4.5),
        "bullet": ParagraphStyle("ResumeBullet", parent=sample["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13.4, textColor=HexColor("#27272a"), leftIndent=12, firstLineIndent=-8, spaceAfter=4),
    }
    lines = _lines(doc)
    story = []
    offset = _resume_header(lines, story, styles, accent) if doc.document_type == "resume" else 0
    for line in lines[offset:]:
        if _is_heading(line):
            story.append(Paragraph(escape(line.upper()), styles["heading"]))
        elif _is_bullet(line):
            story.append(Paragraph(f"&#8226;&nbsp;&nbsp;{escape(_bullet_text(line))}", styles["bullet"]))
        else:
            story.append(Paragraph(escape(line), styles["body"]))

    def footer(canvas, page_doc):
        canvas.saveState()
        canvas.setStrokeColor(HexColor("#e4e4e7"))
        canvas.setLineWidth(0.5)
        canvas.line(18 * mm, 12 * mm, A4[0] - 18 * mm, 12 * mm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(HexColor("#71717a"))
        canvas.drawString(18 * mm, 8 * mm, "CAREERPILOT  ·  RESUME")
        canvas.drawRightString(A4[0] - 18 * mm, 8 * mm, str(page_doc.page))
        canvas.restoreState()

    margin = 20 if template == "minimal" else 18
    SimpleDocTemplate(output, pagesize=A4, leftMargin=margin * mm, rightMargin=margin * mm,
                      topMargin=17 * mm, bottomMargin=18 * mm, title=doc.title,
                      author="CareerPilot").build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_docx_bytes(doc) -> bytes:
    output = Document()
    section = output.sections[0]
    template = doc.template_key or "ats"
    accent = DOCX_ACCENTS.get(template, DOCX_ACCENTS["ats"])
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = output.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(39, 39, 42)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("CAREERPILOT  ·  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(113, 113, 122)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    lines = _lines(doc)
    offset = 0
    if doc.document_type == "resume" and lines:
        p = output.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(lines[0])
        r.bold = True
        r.font.name = "Aptos Display"
        r.font.size = Pt(24)
        r.font.color.rgb = RGBColor(17, 24, 39)
        if len(lines) > 1:
            p = output.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(lines[1])
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(82, 82, 91)
        offset = min(2, len(lines))
        if len(lines) > 2 and not _is_heading(lines[2]):
            p = output.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(lines[2])
            r.bold = True
            r.font.color.rgb = accent
            offset = 3

    for line in lines[offset:]:
        if _is_heading(line):
            p = output.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(line.upper())
            r.bold = True
            r.font.name = "Aptos Display"
            r.font.size = Pt(10.5)
            r.font.color.rgb = accent
        elif _is_bullet(line):
            p = output.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.24)
            p.paragraph_format.first_line_indent = Inches(-0.14)
            p.paragraph_format.space_after = Pt(4)
            p.add_run(_bullet_text(line))
        else:
            p = output.add_paragraph(line)
            p.paragraph_format.widow_control = True

    buffer = io.BytesIO()
    output.save(buffer)
    return buffer.getvalue()
